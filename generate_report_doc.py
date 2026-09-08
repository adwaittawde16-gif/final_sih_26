import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def generate_doc(output_filename):
    doc = docx.Document()

    # Set page margins (0.75 in)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Styles & Colors
    NAVY = RGBColor(15, 23, 42)      # #0f172a
    BLUE = RGBColor(37, 99, 235)     # #2563eb
    DARK_GRAY = RGBColor(51, 65, 85) # #334155
    GRAY = RGBColor(100, 116, 139)   # #64748b

    # Document Header Title Block
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_dept = title_p.add_run("BRIHANMUMBAI POLICE DEPARTMENT — SPECIAL CRIME ANALYSIS UNIT\n")
    run_dept.font.name = "Arial"
    run_dept.font.size = Pt(10)
    run_dept.font.bold = True
    run_dept.font.color.rgb = BLUE

    run_main = title_p.add_run("AI-Powered Criminal Network & Tactical Intelligence System (SIH 26)\n")
    run_main.font.name = "Arial"
    run_main.font.size = Pt(18)
    run_main.font.bold = True
    run_main.font.color.rgb = NAVY

    run_sub = title_p.add_run("Comprehensive Technical Architecture, Module Integration & UI Transformation Report")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = GRAY

    doc.add_paragraph() # Spacer

    # 1. Executive Summary
    h1 = doc.add_heading("1. Executive Summary", level=1)
    h1.runs[0].font.color.rgb = NAVY

    p1 = doc.add_paragraph(
        "This report documents the end-to-end development, full-stack REST API architecture, and UI/UX modernization "
        "of the AI-Powered Criminal Network Analysis System (SIH 26) built for the Brihanmumbai Police Department. "
        "The system fuses multi-source intelligence records—including First Information Reports (FIRs), Call Detail Records (CDRs), "
        "municipal CCTV facial recognition sightings, financial transaction trails, field surveillance observations, and cross-platform "
        "digital footprints—into an interactive, real-time command center interface."
    )
    p1.runs[0].font.name = "Arial"
    p1.runs[0].font.size = Pt(10.5)

    # 2. Detailed Breakdown of 9 Core Intelligence Modules
    h2 = doc.add_heading("2. Full Breakdown of Core Intelligence Modules", level=1)
    h2.runs[0].font.color.rgb = NAVY

    modules_info = [
        ("Command Center Dashboard (/) ", "Unified executive overview presenting system health, priority threat rankings, real-time police alert feeds, and quick module navigation."),
        ("Module 1 — Suspect Threat Index (/threat)", "Algorithmic multi-criteria risk scoring (0-100) combining CCTV, CDR, FIR, Financial, and Surveillance indicators. Includes live weight simulation sliders."),
        ("Module 2 — CDR Call Network Graph (/cdr)", "Interactive D3.js force-directed network graph visualizing call volume, duration, nocturnal frequency, and influencer degree/betweenness centrality."),
        ("Module 3 — CCTV Physical Co-Location (/cctv)", "Spatiotemporal camera sighting correlation detecting physical co-location encounters, facial match confidence, distance proximity, and GPS Leaflet visualizer."),
        ("Module 4 — Gangs & Crime Syndicates (/gangs)", "Graph community detection (Louvain algorithm) automatically clustering suspect entities into named gangs (Gang 1, Gang 2). Includes investigator Confirm, Rename, Merge, and Subgraph tools."),
        ("Module 5 — Financial Intelligence (/financial)", "Suspicious UPI and ATM transaction analysis tracking money laundering trails, peer transfers, failed ATM withdrawals, and high-risk expenditures."),
        ("Module 6 — Nocturnal Call Anomalies (/nocturnal)", "Midnight communication monitoring (00:00–06:00 IST) flagging suspicious call frequencies and cell tower handover hotspots with Leaflet GPS mapping."),
        ("Module 7 — Field Surveillance Density (/surveillance)", "Special Branch field officer observation logs, panchnama documentation, and spatial patrol spot coordinates."),
        ("Module 8 — 360° Suspect Dossiers & NLP (/dossiers)", "Cross-database search, 360° markdown dossier viewer, Interactive Forensic Timeline (color-coded FIR, CDR, CCTV, Financial events), and live NLP FIR Entity & Co-Accused Extractor tool."),
        ("Module 9 — Digital Footprint & Social Media (/social-media)", "Cross-platform login location co-occurrence analysis, handle surveillance, sentiment classification, and flagged post monitoring.")
    ]

    for title, desc in modules_info:
        p = doc.add_paragraph()
        r_t = p.add_run(f"• {title}: ")
        r_t.font.bold = True
        r_t.font.color.rgb = BLUE
        r_t.font.name = "Arial"
        r_t.font.size = Pt(10)
        r_d = p.add_run(desc)
        r_d.font.name = "Arial"
        r_d.font.size = Pt(10)

    doc.add_paragraph()

    # 3. UI Transformation & Key Differences from Original GitHub Repo
    h3 = doc.add_heading("3. Detailed Comparison & UI Transformation vs. Original Repo", level=1)
    h3.runs[0].font.color.rgb = NAVY

    p_diff = doc.add_paragraph(
        "The current platform represents a major architectural and visual evolution from the original GitHub repository "
        "(adwaittawde16-gif/new_sih). The original project consisted primarily of disconnected Python CLI scripts generating "
        "static HTML files. Below is a comprehensive comparison detailing all architectural and UI/UX enhancements:"
    )
    p_diff.runs[0].font.name = "Arial"
    p_diff.runs[0].font.size = Pt(10.5)

    # Comparison Table
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    hdr_cells = table.rows[0].cells
    headers = ["Feature / Dimension", "Original GitHub Repo", "Our Modernized Platform"]
    widths = [Inches(1.8), Inches(2.3), Inches(2.9)]

    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        hdr_cells[i].width = widths[i]
        set_cell_background(hdr_cells[i], "0F172A")
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.name = "Arial"
            run.font.size = Pt(9.5)

    diff_data = [
        ("Architecture & Delivery", "Disconnected scripts outputting static .html files (timeline_*.html, cctv_co_location_map.html).", "Unified Full-Stack App: FastAPI REST JSON Backend + Next.js 14 React SPA (0 static HTML files/iframes)."),
        ("UI Aesthetic & Theme", "Dark Streamlit/raw browser HTML with generic table layouts.", "Light Mode Command Center: Tailored slate-900 typography, custom color badges, glassmorphism cards, and subtle micro-animations."),
        ("Gang Detection & Management", "Basic ring leader text list; no automated community clustering.", "Automated Louvain Clustering: Auto-detects Gang 1, Gang 2 with investigator Confirm, Rename, Merge, and Subgraph tools."),
        ("Forensic Timeline", "Static HTML generator per suspect executed via CLI.", "Native React Component (<ForensicTimeline />): Color-coded filters (FIR, CDR, CCTV, Financial) with expandable JSON metadata."),
        ("Social Media & Footprint", "Console printouts & isolated HTML generator.", "Dedicated Module 9 (/social-media): Login location co-occurrence table, monitored handles, and sentiment analysis."),
        ("NLP FIR Entity Extraction", "Terminal-only script (nlp_fir_analyzer.py).", "Interactive Web Parser (<FIRParserTool />): Paste raw FIR narratives directly into UI for instant suspect, co-accused, location & IPC extraction."),
        ("Geospatial Mapping", "Static Folium HTML file output.", "Shared <LeafletMapView /> Component: OpenStreetMap tile rendering integrated across /cctv, /nocturnal, and /surveillance."),
        ("Terminology & Labeling", "Generic labels like 'Threat Score Leaderboard'.", "Tactical Police Terminology: Replaced with 'Suspect Threat Index' and 'Priority Risk Rankings'.")
    ]

    for row_idx, (dim, old_val, new_val) in enumerate(diff_data):
        row_cells = table.add_row().cells
        bg_color = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"

        row_cells[0].text = dim
        row_cells[1].text = old_val
        row_cells[2].text = new_val

        for c_idx, cell in enumerate(row_cells):
            cell.width = widths[c_idx]
            set_cell_background(cell, bg_color)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = "Arial"
                run.font.size = Pt(9)
                if c_idx == 0:
                    run.font.bold = True
                    run.font.color.rgb = NAVY
                elif c_idx == 1:
                    run.font.color.rgb = DARK_GRAY
                else:
                    run.font.color.rgb = BLUE

    doc.add_paragraph() # Spacer

    # 4. Technical Stack Summary
    h4 = doc.add_heading("4. Technical Stack & Deployment", level=1)
    h4.runs[0].font.color.rgb = NAVY

    p_stack = doc.add_paragraph(
        "• Backend Framework: FastAPI (Python 3.10) with Pydantic v2 schemas and modular APIRouters.\n"
        "• Data Analytics Engine: Pandas, NetworkX, IntelligenceEngine in-memory data structures.\n"
        "• Frontend Framework: Next.js 14 (App Router), TypeScript, React 18, Tailwind CSS.\n"
        "• Visualization Libraries: D3.js (Force Network Graphs), Leaflet.js & OpenStreetMap (GPS Maps), Lucide Icons.\n"
        "• Verification: Clean production build verified via 'npm run build' with zero compiler errors across all 14 routes."
    )
    p_stack.runs[0].font.name = "Arial"
    p_stack.runs[0].font.size = Pt(10)

    # Save document
    doc.save(output_filename)
    print(f"[OK] Successfully generated Word Document: {output_filename}")

if __name__ == "__main__":
    generate_doc("SIH_26_Police_Intelligence_System_Report.docx")
