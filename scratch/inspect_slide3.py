import os, pptx
from pptx.util import Inches

ppt_path = os.path.abspath('SIH2026Presentation-2_Updated.pptx')
if not os.path.exists(ppt_path):
    ppt_path = r'C:\Users\Bhavya Parekh\AppData\Local\Packages\5319275A.WhatsAppDesktop_cv1g1gvanyjgm\LocalState\sessions\B196F22353273CED48C721075BF4CF89D3B24CD9\transfers\2026-38\SIH2026Presentation-2.pptx'

prs = pptx.Presentation(ppt_path)
slide3 = prs.slides[2]

print(f"Slide 3 Total Shapes: {len(slide3.shapes)}")
for i, shape in enumerate(slide3.shapes):
    txt = shape.text[:80].replace('\n', ' ') if shape.has_text_frame else '[NON-TEXT / SHAPE / IMAGE]'
    left = shape.left/Inches(1) if shape.left else 0
    top = shape.top/Inches(1) if shape.top else 0
    w = shape.width/Inches(1) if shape.width else 0
    h = shape.height/Inches(1) if shape.height else 0
    print(f"Shape {i+1}: pos=({left:.2f}, {top:.2f}), size=({w:.2f}x{h:.2f}) -> {txt}")
