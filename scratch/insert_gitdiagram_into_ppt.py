import os, pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

ppt_path = r'C:\Users\Bhavya Parekh\AppData\Local\Packages\5319275A.WhatsAppDesktop_cv1g1gvanyjgm\LocalState\sessions\B196F22353273CED48C721075BF4CF89D3B24CD9\transfers\2026-38\SIH2026Presentation-2.pptx'

print(f"Loading presentation from: {ppt_path}")
prs = pptx.Presentation(ppt_path)
slide3 = prs.slides[2]

DARK_BLUE = RGBColor(30, 58, 138)

# Delete non-header shapes on Slide 3
shapes_to_remove = []
for shape in slide3.shapes:
    if shape.has_text_frame:
        txt = shape.text.strip()
        if "TECHNICAL APPROACH" in txt or "SYSTEM ARCHITECTURE" in txt:
            shape.text_frame.clear()
            p = shape.text_frame.paragraphs[0]
            p.text = "SYSTEM ARCHITECTURE & REPOSITORY FLOWCHART"
            p.font.size = Pt(22)
            p.font.bold = True
            p.font.color.rgb = DARK_BLUE
            shape.left = Inches(0.5)
            shape.top = Inches(0.25)
            shape.width = Inches(12.33)
            shape.height = Inches(0.7)
        elif "Ctrl Innovate" in txt or "@SIH Idea submission" in txt or txt == "3" or txt == "":
            continue
        else:
            shapes_to_remove.append(shape)
    else:
        shapes_to_remove.append(shape)

# Remove background/content shapes
spTree = slide3.shapes._spTree
for shp in shapes_to_remove:
    try:
        spTree.remove(shp._element)
    except Exception as e:
        pass

# Insert gitdiagram architecture flowchart image
gitdiag_img_path = os.path.abspath("frontend/public/gitdiagram_architecture.png")
if os.path.exists(gitdiag_img_path):
    slide3.shapes.add_picture(gitdiag_img_path, Inches(0.4), Inches(1.0), width=Inches(12.5), height=Inches(5.7))
    print("Inserted GitDiagram Architecture Image into Slide 3 cleanly!")

out_final = os.path.abspath("SIH2026Presentation_GitDiagram.pptx")
prs.save(out_final)
print(f"Presentation saved successfully to:\n  - {out_final}")
