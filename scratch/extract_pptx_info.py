import os
import pptx
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

pptx_path = r"C:\Users\Bhavya Parekh\AppData\Local\Packages\5319275A.WhatsAppDesktop_cv1g1gvanyjgm\LocalState\sessions\B196F22353273CED48C721075BF4CF89D3B24CD9\transfers\2026-38\SIH2026Presentation-2.pptx"

prs = Presentation(pptx_path)
print(f"Loaded PPTX with {len(prs.slides)} slides.")

output_dir = r"c:\Users\Bhavya Parekh\OneDrive\Pictures\Desktop\SIH 26\final final\final_sih_26\frontend\public"
os.makedirs(output_dir, exist_ok=True)

img_count = 0
for slide_idx, slide in enumerate(prs.slides):
    print(f"\n--- Slide {slide_idx + 1} ---")
    for shape in slide.shapes:
        if shape.has_text_frame:
            text = shape.text.strip().replace("\n", " ")
            if text:
                print(f"  [Text]: {text[:120].encode('ascii', 'replace').decode()}")
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            image = shape.image
            image_bytes = image.blob
            image_filename = f"pptx_slide_{slide_idx + 1}_img_{img_count + 1}.{image.ext}"
            image_path = os.path.join(output_dir, image_filename)
            with open(image_path, "wb") as f:
                f.write(image_bytes)
            print(f"  [Extracted Image Saved]: {image_filename}")
            img_count += 1
