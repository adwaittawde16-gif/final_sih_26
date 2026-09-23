import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'  <w:left w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def create_document():
    doc = Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styling helper colors
    NAVY = RGBColor(15, 23, 42)      # #0F172A
    BLUE = RGBColor(29, 78, 216)     # #1D4ED8
    DARK_BLUE = RGBColor(30, 58, 138) # #1E3A8A
    GRAY = RGBColor(100, 116, 139)   # #64748B
    DARK_GRAY = RGBColor(51, 65, 85) # #334155
    BLACK = RGBColor(15, 23, 42)

    # Document Header Title Block
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = title_p.add_run("TACTICAL INTELLIGENCE COMMAND CENTER\n")
    run_sub.font.name = 'Arial'
    run_sub.font.size = Pt(11)
    run_sub.font.bold = True
    run_sub.font.color.rgb = BLUE

    run_title = title_p.add_run("PROJECT FLOWCHART & SYSTEM ARCHITECTURE GUIDE\n")
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = NAVY

    run_desc = title_p.add_run("Problem Analysis, Proposed Solution, Layered System Architecture & Flowchart Breakdown")
    run_desc.font.name = 'Calibri'
    run_desc.font.size = Pt(12)
    run_desc.font.italic = True
    run_desc.font.color.rgb = GRAY

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Callout Box: Executive Overview
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F0F4FF")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    p = cell.paragraphs[0]
    r_hdr = p.add_run("📌 EXECUTIVE SUMMARY\n")
    r_hdr.font.name = 'Arial'
    r_hdr.font.size = Pt(12)
    r_hdr.font.bold = True
    r_hdr.font.color.rgb = DARK_BLUE
    
    r_body = p.add_run(
        "The Tactical Intelligence Command & Evidence Attribution Engine is an enterprise law-enforcement platform "
        "designed to solve critical bottlenecks in police criminal investigations. By unifying disconnected streams "
        "(CDR telecom records, CCTV visual surveillance, FIR police complaints, financial logs, and social media networks) "
        "into a single 100-Point Dynamic Threat Classifier and Explainable AI (XAI) engine, the system delivers court-admissible "
        "mathematical proofs, counterfactual evidence lineage, and spatiotemporal tracking for high-priority syndicates."
    )
    r_body.font.name = 'Calibri'
    r_body.font.size = Pt(11)
    r_body.font.color.rgb = DARK_GRAY

    doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # Section 1: Problem Statement
    h1 = doc.add_heading(level=1)
    r = h1.add_run("1. PROBLEM STATEMENT: CHALLENGES IN MODERN LAW ENFORCEMENT")
    r.font.name = 'Arial'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = NAVY

    problems = [
        ("Data Fragmentation & Intelligence Silos", 
         "Police departments process huge volumes of data across completely isolated formats: telecom Call Detail Records (CDRs), CCTV camera feeds, PDF FIR complaint archives, bank transaction ledgers, and social media footprints. Without automated correlation, cross-referencing suspects across jurisdictions takes weeks or months."),
        
        ("Manual Analysis & Information Overload", 
         "Investigating officers manually compare thousands of phone tower ping logs and video recordings. High-priority criminal syndicates exploit these delay windows to change locations, discard SIM cards, or obscure evidence trail."),
        
        ("Lack of Legal Admissibility & 'Black-Box' AI Risks", 
         "Standard machine learning models generate threat alerts without transparent mathematical justifications. Courts (under statutory frameworks like MCOCA / IPC Sec 120B) reject 'black-box' predictions unless investigators present clear, step-by-step feature attribution and evidence lineage."),
        
        ("Evidence Tampering & Chain-of-Custody Gaps", 
         "Digital evidence logs stored in conventional spreadsheets or databases lack cryptographic protection, leaving them vulnerable to defense challenges regarding authenticity, modifications, or chain-of-custody breaks.")
    ]

    for title, desc in problems:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(6)
        r_t = p.add_run(f"• {title}: ")
        r_t.font.name = 'Arial'
        r_t.font.size = Pt(11)
        r_t.font.bold = True
        r_t.font.color.rgb = DARK_BLUE
        
        r_d = p.add_run(desc)
        r_d.font.name = 'Calibri'
        r_d.font.size = Pt(11)
        r_d.font.color.rgb = DARK_GRAY

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 2: Solution Provided
    h2 = doc.add_heading(level=1)
    r = h2.add_run("2. SOLUTION PROVIDED: TACTICAL INTELLIGENCE PLATFORM")
    r.font.name = 'Arial'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = NAVY

    solutions = [
        ("Multi-Source Automated Ingestion Engine", 
         "In-built parsers digest CDR telecom logs, FIR text PDFs, CCTV video streams, bank financial files, and social media handles into a standardized relational and graph model."),
        
        ("100-Point Dynamic Threat Indexing System", 
         "An algorithmic scoring engine evaluates suspects across 6 quantitative crime vectors (CCTV Sighting, Nocturnal CDR Interceptions, FIR Severity, Criminal History, Financial Risk, and Social Media Association) to calculate a unified 0-100 Threat Score with color-coded risk tiers (Critical Red, Elevated Orange, Standard Blue)."),
        
        ("Explainable AI (XAI) & Counterfactual Evidence Attribution", 
         "Breakdown of exact mathematical contributions (+28.6 pts CCTV, +20 pts CDR, +15 pts FIR) paired with interactive counterfactual simulation ('What if phone was switched off at 02:15 AM?') to generate court-proof prosecution dossiers."),
        
        ("Spatiotemporal Timeline & CCTV Playback", 
         "Interactive video player synchronized with cell-tower location hops, placing visual face-match captures on geographical maps in real time."),
        
        ("Graph Centrality & Syndicate Discovery", 
         "Automated network analysis algorithms (Degree Centrality, Betweenness, Gephi GEXF exports) identify hidden gang leaders, kingpins, and intermediary contacts."),
        
        ("Tamper-Evident Cryptographic SHA-256 Audit Trail", 
         "Every search, record update, and AI risk prediction is hashed sequentially into a SHA-256 forensic log ledger, guaranteeing unalterable chain-of-custody for judicial scrutiny.")
    ]

    for title, desc in solutions:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(6)
        r_t = p.add_run(f"✔ {title}: ")
        r_t.font.name = 'Arial'
        r_t.font.size = Pt(11)
        r_t.font.bold = True
        r_t.font.color.rgb = BLUE
        
        r_d = p.add_run(desc)
        r_d.font.name = 'Calibri'
        r_d.font.size = Pt(11)
        r_d.font.color.rgb = DARK_GRAY

    doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # Section 3: High-Level System Flowchart Diagram
    h3 = doc.add_heading(level=1)
    r = h3.add_run("3. SYSTEM FLOWCHART (END-TO-END PIPELINE)")
    r.font.name = 'Arial'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = NAVY

    # Add visual image from PPT
    import os
    ppt_img_path = os.path.abspath("frontend/public/image2.png")
    if os.path.exists(ppt_img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(8)
        p_img.add_run().add_picture(ppt_img_path, width=Inches(6.5))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run("Figure 1: High-Resolution Tactical System Architecture & Core Workflow Diagram (from SIH2026Presentation-2.pptx)")
        r_cap.font.name = 'Calibri'
        r_cap.font.size = Pt(9.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = GRAY

    # Flowchart Visual Box (Text Box Table)
    tbl_flow = doc.add_table(rows=1, cols=1)
    tbl_flow.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_flow = tbl_flow.cell(0, 0)
    set_cell_background(cell_flow, "F8FAFC")
    set_cell_margins(cell_flow, top=140, bottom=140, left=150, right=150)
    
    p_flow = cell_flow.paragraphs[0]
    p_flow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    flowchart_ascii = (
        "+-----------------------------------------------------------------------------------+\n"
        "|                             1. MULTI-SOURCE DATA INGESTION                         |\n"
        "|  [CDR Phone Logs]   [CCTV Cameras]   [FIR Complaints]   [Bank Logs]   [Social Media]  |\n"
        "+-----------------------------------------------------------------------------------+\n"
        "                                          │                                          \n"
        "                                          ▼                                          \n"
        "+-----------------------------------------------------------------------------------+\n"
        "|                        2. FEATURE EXTRACTION & NLP PREPROCESSING                  |\n"
        "|  • Regex CDR Parser   • Facial Recog (OpenCV)   • SpaCy Entity Extraction (FIR)   |\n"
        "+-----------------------------------------------------------------------------------+\n"
        "                                          │                                          \n"
        "                                          ▼                                          \n"
        "+-----------------------------------------------------------------------------------+\n"
        "|                      3. GRAPH ANALYTICS & ANOMALY CORRELATION                     |\n"
        "|  • Network Centrality (NetworkX)   • Nocturnal Call Spikes   • Co-Location Clustering |\n"
        "+-----------------------------------------------------------------------------------+\n"
        "                                          │                                          \n"
        "                                          ▼                                          \n"
        "+-----------------------------------------------------------------------------------+\n"
        "|                   4. THREAT SCORING & EXPLAINABLE AI (XAI) MATRIX                 |\n"
        "|  • 100-Point Threat Score   • Mathematical Attribution   • Counterfactual Simulation  |\n"
        "+-----------------------------------------------------------------------------------+\n"
        "                                          │                                          \n"
        "                                          ▼                                          \n"
        "+-----------------------------------------------------------------------------------+\n"
        "|                   5. TACTICAL COMMAND DASHBOARD & EVIDENCE EXPORT                  |\n"
        "|  • Live Maps & CCTV Timeline   • Audio AI Copilot   • Court Dossier PDF & SHA Audit|\n"
        "+-----------------------------------------------------------------------------------+"
    )
    
    r_ascii = p_flow.add_run(flowchart_ascii)
    r_ascii.font.name = 'Consolas'
    r_ascii.font.size = Pt(8.5)
    r_ascii.font.bold = True
    r_ascii.font.color.rgb = DARK_BLUE

    doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # Section 4: Detailed Flowchart Breakdown Table ("HOW IT WORKS")
    h4 = doc.add_heading(level=1)
    r = h4.add_run("4. DETAILED STEP-BY-STEP PIPELINE BREAKDOWN ('HOW IT WORKS')")
    r.font.name = 'Arial'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = NAVY

    table_data = [
        ("Step 1: Multi-Source Ingestion", 
         "Input Data Stream", 
         "Raw CSV files, JSON payloads, RTSP camera streams, and PDF police records are fed into the ingestion service."),
        
        ("Step 2: Parsing & Entity Extraction", 
         "NLP & Computer Vision", 
         "SpaCy extracts criminal names, addresses, and IPC sections from FIRs. Facial recognition scans CCTV frames. Regex parses cell tower IDs & call durations."),
        
        ("Step 3: Graph & Spatial Correlation", 
         "Correlation Engine", 
         "Nodes (suspects) and edges (calls/co-locations) are stored in graph memory. Spatiotemporal algorithms detect simultaneous presence of suspects near crime scenes."),
        
        ("Step 4: Threat Indexing & Scoring", 
         "Dynamic Classifier", 
         "Calculates composite 0-100 Threat Score based on 6 weighted vectors. Subjects scoring >= 75 are flagged as Critical / Level-1 Red."),
        
        ("Step 5: Explainable AI Attribution", 
         "SHAP/LIME Evidence Engine", 
         "Generates percentage influence breakdown (+28.6% CCTV, +20.0% CDR) and executes counterfactual scenarios for statutory chargeability (MCOCA/IPC 120B)."),
        
        ("Step 6: Dashboard Command Presentation", 
         "Next.js Web UI", 
         "Displays subject dossiers, interactive GEXF network graphs, spatiotemporal CCTV timelines, and voice-assisted copilot AI."),
        
        ("Step 7: Security & Audit Logging", 
         "SHA-256 Forensic Audit", 
         "All queries, user edits, and export actions produce cryptographically chained SHA-256 hashes to guarantee court tamper-evidence.")
    ]

    t = doc.add_table(rows=len(table_data) + 1, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t)

    # Table Header
    hdr_cells = t.rows[0].cells
    hdr_titles = ["Pipeline Phase", "Core Engine / Method", "Operational Function & Output"]
    for i, text in enumerate(hdr_titles):
        set_cell_background(hdr_cells[i], "1E3A8A")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        r = p.add_run(text)
        r.font.name = 'Arial'
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    # Table Rows
    for row_idx, data in enumerate(table_data, start=1):
        row_cells = t.rows[row_idx].cells
        bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = row_cells[col_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.name = 'Calibri'
            r.font.size = Pt(9.5)
            if col_idx == 0:
                r.font.bold = True
                r.font.color.rgb = DARK_BLUE
            else:
                r.font.color.rgb = DARK_GRAY

    doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # Section 5: Comparative Impact Matrix
    h5 = doc.add_heading(level=1)
    r = h5.add_run("5. COMPARATIVE MATRIX: TRADITIONAL VS. TACTICAL COMMAND PLATFORM")
    r.font.name = 'Arial'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = NAVY

    comp_data = [
        ("Investigation Speed", "Manual paper & spreadsheet check (Weeks)", "Automated multi-stream correlation (Seconds)"),
        ("Cross-Source Correlation", "Disjointed analysis across departments", "Unified Graph & Spatiotemporal Engine"),
        ("AI Trust & Court Proof", "Black-box AI rejected in cross-examination", "Explainable AI (XAI) feature attribution"),
        ("Evidence Tamper Proof", "Standard database logs prone to edits", "Cryptographic SHA-256 immutable audit chain"),
        ("Syndicate Detection", "Only direct contacts identified", "Graph Centrality algorithms uncover hidden leaders")
    ]

    t_comp = doc.add_table(rows=len(comp_data) + 1, cols=3)
    t_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_comp)

    hdr_cells_c = t_comp.rows[0].cells
    hdr_titles_c = ["Dimension", "Traditional Police Workflow", "Tactical Intelligence Command Engine"]
    for i, text in enumerate(hdr_titles_c):
        set_cell_background(hdr_cells_c[i], "0F172A")
        set_cell_margins(hdr_cells_c[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells_c[i].paragraphs[0]
        r = p.add_run(text)
        r.font.name = 'Arial'
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    for row_idx, data in enumerate(comp_data, start=1):
        row_cells = t_comp.rows[row_idx].cells
        bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = row_cells[col_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.name = 'Calibri'
            r.font.size = Pt(9.5)
            if col_idx == 0:
                r.font.bold = True
                r.font.color.rgb = NAVY
            elif col_idx == 1:
                r.font.color.rgb = RGBColor(185, 28, 28) # Dark Red
            else:
                r.font.bold = True
                r.font.color.rgb = RGBColor(21, 128, 61) # Dark Green

    doc.add_paragraph().paragraph_format.space_after = Pt(20)

    # Footer note
    p_ftr = doc.add_paragraph()
    p_ftr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_ftr = p_ftr.add_run("Generated for Brihanmumbai Police / Special Crime Analysis Unit • Confidential")
    r_ftr.font.name = 'Calibri'
    r_ftr.font.size = Pt(9)
    r_ftr.font.italic = True
    r_ftr.font.color.rgb = GRAY

    output_filename = "SIH_26_Project_Flowchart_and_Solution_Guide.docx"
    doc.save(output_filename)
    print(f"Document successfully created: {output_filename}")

if __name__ == "__main__":
    create_document()
