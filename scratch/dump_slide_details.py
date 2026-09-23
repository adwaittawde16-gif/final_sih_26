import pptx
from pptx.util import Inches

ppt_path = r'C:\Users\Bhavya Parekh\AppData\Local\Packages\5319275A.WhatsAppDesktop_cv1g1gvanyjgm\LocalState\sessions\B196F22353273CED48C721075BF4CF89D3B24CD9\transfers\2026-38\SIH2026Presentation-2.pptx'
prs = pptx.Presentation(ppt_path)

for i, slide in enumerate(prs.slides):
    print(f"\n==================== SLIDE {i+1} ====================")
    for j, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            left = shape.left/Inches(1) if shape.left else 0
            top = shape.top/Inches(1) if shape.top else 0
            w = shape.width/Inches(1) if shape.width else 0
            h = shape.height/Inches(1) if shape.height else 0
            print(f"\n  --- Shape {j+1} [pos: ({left:.2f}, {top:.2f}), size: ({w:.2f}x{h:.2f})] ---")
            for p in shape.text_frame.paragraphs:
                p_text = p.text.encode('ascii', errors='ignore').decode('ascii')
                font_sz = p.font.size.pt if p.font and p.font.size else 'default'
                print(f"    P (font={font_sz}): {p_text}")
