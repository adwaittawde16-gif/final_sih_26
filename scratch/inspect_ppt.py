import pptx
from pptx.util import Inches

ppt_path = r'C:\Users\Bhavya Parekh\AppData\Local\Packages\5319275A.WhatsAppDesktop_cv1g1gvanyjgm\LocalState\sessions\B196F22353273CED48C721075BF4CF89D3B24CD9\transfers\2026-38\SIH2026Presentation-2.pptx'
prs = pptx.Presentation(ppt_path)

print(f"Slide dimensions: width={prs.slide_width/Inches(1):.2f}\", height={prs.slide_height/Inches(1):.2f}\"")
for i, slide in enumerate(prs.slides):
    print(f"\n=== SLIDE {i+1} (shapes: {len(slide.shapes)}) ===")
    for j, shape in enumerate(slide.shapes):
        left = shape.left/Inches(1) if shape.left else 0
        top = shape.top/Inches(1) if shape.top else 0
        width = shape.width/Inches(1) if shape.width else 0
        height = shape.height/Inches(1) if shape.height else 0
        text = shape.text[:100].replace('\n', ' ') if shape.has_text_frame else '[NO TEXT]'
        safe_text = text.encode('ascii', errors='ignore').decode('ascii')
        print(f"  Shape {j+1}: pos=({left:.2f}, {top:.2f}), size=({width:.2f}x{height:.2f}) -> {safe_text}")
