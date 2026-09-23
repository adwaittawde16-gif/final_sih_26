import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_sih_6slide_presentation(output_path):
    prs = Presentation()
    # 16:9 Widescreen (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette - Professional Dark Mode Command Center Theme
    COLOR_BG = RGBColor(15, 23, 42)          # Deep Navy Slate (#0F172A)
    COLOR_CARD = RGBColor(30, 41, 59)        # Card Background (#1E293B)
    COLOR_BORDER = RGBColor(51, 65, 85)      # Border Grey (#334155)
    COLOR_CYAN = RGBColor(56, 189, 248)      # Tech Cyan Accent (#38BDF8)
    COLOR_GOLD = RGBColor(245, 158, 11)      # Gold Highlight (#F59E0B)
    COLOR_RED = RGBColor(239, 68, 68)        # Crimson Alert (#EF4444)
    COLOR_TEXT_MAIN = RGBColor(241, 245, 249)# Off White (#F1F5F9)
    COLOR_TEXT_MUTED = RGBColor(148, 163, 184)# Muted Slate (#94A3B8)
    COLOR_GREEN = RGBColor(34, 197, 94)      # Success Green (#22C55E)

    def set_slide_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = COLOR_BG

    def add_header(slide, slide_num_str, title_text, category_text):
        # Top banner category
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(10), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_CYAN
        p_cat.font.name = 'Calibri'

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(10), Inches(0.75))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_TEXT_MAIN
        p_title.font.name = 'Calibri'

        # Slide Number Badge
        num_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(11.8), Inches(0.45), Inches(0.8), Inches(0.4)
        )
        num_shape.fill.solid()
        num_shape.fill.fore_color.rgb = COLOR_CARD
        num_shape.line.color.rgb = COLOR_CYAN
        num_shape.line.width = Pt(1)
        tf_num = num_shape.text_frame
        p_num = tf_num.paragraphs[0]
        p_num.text = slide_num_str
        p_num.alignment = PP_ALIGN.CENTER
        p_num.font.size = Pt(12)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_CYAN

        # Bottom Footer line
        footer_line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(7.0), Inches(11.833), Inches(0.02)
        )
        footer_line.fill.solid()
        footer_line.fill.fore_color.rgb = COLOR_BORDER
        footer_line.line.fill.background()

        # Footer Text
        foot_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.833), Inches(0.3))
        tf_foot = foot_box.text_frame
        p_foot = tf_foot.paragraphs[0]
        p_foot.text = "Smart India Hackathon 2026 | Official 6-Slide Submission Presentation | Team CTRL INNOVATE"
        p_foot.font.size = Pt(9)
        p_foot.font.color.rgb = COLOR_TEXT_MUTED
        p_foot.font.name = 'Calibri'

    blank_layout = prs.slide_layouts[6]

    # ==========================================
    # SLIDE 1: TITLE SLIDE (Official SIH Format)
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)

    top_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.15))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = COLOR_CYAN
    top_bar.line.fill.background()

    badge = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.6), Inches(3.4), Inches(0.4))
    badge.fill.solid()
    badge.fill.fore_color.rgb = COLOR_CARD
    badge.line.color.rgb = COLOR_CYAN
    tf_b = badge.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.text = "SMART INDIA HACKATHON 2026"
    p_b.font.size = Pt(11)
    p_b.font.bold = True
    p_b.font.color.rgb = COLOR_GOLD
    p_b.alignment = PP_ALIGN.CENTER

    t_box = slide1.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.7), Inches(2.2))
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    
    p1 = tf_t.paragraphs[0]
    p1.text = "AI-Powered Criminal Network & Tactical Intelligence System"
    p1.font.size = Pt(30)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_TEXT_MAIN

    p2 = tf_t.add_paragraph()
    p2.text = "Automated Call Detail Record Graph Mining, CCTV Co-Location Verification & SHA-256 Tamper-Evident Audit Command Center"
    p2.font.size = Pt(14)
    p2.font.color.rgb = COLOR_CYAN
    p2.space_before = Pt(8)

    # 2 Column Metadata Cards
    left_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.7), Inches(5.6), Inches(3.0))
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = COLOR_CARD
    left_card.line.color.rgb = COLOR_BORDER
    tf_lc = left_card.text_frame
    tf_lc.word_wrap = True
    
    p = tf_lc.paragraphs[0]
    p.text = "PROBLEM STATEMENT METADATA"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_GOLD
    p.space_after = Pt(6)

    meta_items_left = [
        ("Problem Statement ID:", "SIH1682 / POLICE-INTEL-2026"),
        ("Problem Statement Title:", "AI-Driven Crime Ring Mapping & Multi-Modal Threat Analytics"),
        ("PS Category:", "Software / Security & Law Enforcement"),
        ("Domain Bucket:", "Cyber Security, AI & National Security"),
        ("Ministry / Organization:", "Brihanmumbai Police Dept / Ministry of Home Affairs")
    ]
    for lbl, val in meta_items_left:
        p = tf_lc.add_paragraph()
        p.text = f"• {lbl} "
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_CYAN
        run = p.add_run()
        run.text = val
        run.font.bold = False
        run.font.color.rgb = COLOR_TEXT_MAIN

    right_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(3.7), Inches(5.7), Inches(3.0))
    right_card.fill.solid()
    right_card.fill.fore_color.rgb = COLOR_CARD
    right_card.line.color.rgb = COLOR_BORDER
    tf_rc = right_card.text_frame
    tf_rc.word_wrap = True

    p = tf_rc.paragraphs[0]
    p.text = "TEAM & SUBMISSION DETAILS"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_GOLD
    p.space_after = Pt(6)

    meta_items_right = [
        ("Team Name:", "CTRL INNOVATE"),
        ("Team Leader:", "Bhavya Parekh (Lead Architect & Fullstack Dev)"),
        ("Team Members:", "Special Crime Analytics & AI Development Unit"),
        ("Target End-Users:", "Police Special Crime Branch & Field Command Vans"),
        ("Project Status:", "Fully Operational Functional Prototype")
    ]
    for lbl, val in meta_items_right:
        p = tf_rc.add_paragraph()
        p.text = f"• {lbl} "
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_CYAN
        run = p.add_run()
        run.text = val
        run.font.bold = False
        run.font.color.rgb = COLOR_TEXT_MAIN

    # ==========================================
    # SLIDE 2: PROPOSED SOLUTION & TECH APPROACH
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_header(slide2, "01/05 [Slide 2]", "Proposed Solution & Detailed Technical Approach", "Core Solution & Innovation")

    # Left Side: The Problem & Solution Pitch
    left_card_s2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.2))
    left_card_s2.fill.solid()
    left_card_s2.fill.fore_color.rgb = COLOR_CARD
    left_card_s2.line.color.rgb = COLOR_RED
    left_card_s2.line.width = Pt(1.5)
    tf_l2 = left_card_s2.text_frame
    tf_l2.word_wrap = True

    p = tf_l2.paragraphs[0]
    p.text = "THE PROBLEM & SOLUTION VISION"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_GOLD
    p.space_after = Pt(6)

    prob_sol_text = [
        ("The Challenge:", "CDRs, CCTV logs, and FIRs sit in disconnected silos. Manual cross-referencing takes weeks during active syndicate investigations, leaving high-risk repeat offenders undetected."),
        ("Our Solution:", "A unified Tactical Command Center that correlates multi-source police data into automated network graphs and dynamic 6-parameter threat ratings."),
        ("Key Innovation:", "Combines call log interaction frequencies with CCTV physical co-location timestamps to confirm physical suspect encounters (e.g. 97% match confidence at MH-CCTV-2466).")
    ]
    for lbl, val in prob_sol_text:
        p = tf_l2.add_paragraph()
        p.text = f"• {lbl} "
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_CYAN
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = val
        run.font.bold = False
        run.font.color.rgb = COLOR_TEXT_MAIN

    # Right Side: 4 Technical Pillars
    pillars_s2 = [
        ("1. Multi-Modal CDR Graph Engine", COLOR_CYAN, "Constructs interactive NetworkX call interaction graphs & detects hidden crime ring syndicates."),
        ("2. Physical CCTV Meeting Miner", COLOR_GOLD, "Verifies frequent caller pairs captured by same CCTV camera within minutes."),
        ("3. 6-Factor Composite Threat Index", COLOR_RED, "Scores suspects (0-100): CCTV (30), CDR (20), FIR (15), History (15), Nocturnal (10), Financials (10)."),
        ("4. SHA-256 Tamper-Evident Audit", COLOR_GREEN, "Hashes every query log entry with prior log hash to guarantee 100% court-admissible audit proof.")
    ]

    for idx, (title, color, desc) in enumerate(pillars_s2):
        y = Inches(1.5 + idx * 1.3)
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), y, Inches(5.7), Inches(1.2))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = color
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = color

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = COLOR_TEXT_MAIN
        p_d.space_before = Pt(2)

    # ==========================================
    # SLIDE 3: TECHNICAL ARCHITECTURE & WORKFLOW
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_header(slide3, "02/05 [Slide 3]", "Technical Architecture & Operational Workflow", "System Architecture & Flow")

    # 3 Architecture Tiers on Top
    tiers = [
        ("FRONTEND DASHBOARD", COLOR_CYAN, "Next.js 14 App Router, React 18, TypeScript, D3.js Force Network Graphs, Leaflet GPS Maps."),
        ("REST BACKEND & ENGINE", COLOR_GOLD, "FastAPI REST Server, Pandas & NetworkX graph engine, 6-Factor Threat Classifier, Regex FIR NLP Parser."),
        ("DATA & SECURITY LAYER", COLOR_GREEN, "Structured CSV/JSON Datasets, SQLite SHA-256 Tamper-Evident Hash Audit Chain, Non-root Docker (UID 1000).")
    ]

    for idx, (title, color, desc) in enumerate(tiers):
        x = Inches(0.8 + idx * 3.95)
        card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.5), Inches(3.7), Inches(1.5))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = color
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = color

        p_sub = tf.add_paragraph()
        p_sub.text = desc
        p_sub.font.size = Pt(10)
        p_sub.font.color.rgb = COLOR_TEXT_MAIN
        p_sub.space_before = Pt(4)

    # 5 Operational Workflow Stages on Bottom
    stages = [
        ("1. INGESTION", "Ingests CDR logs, CCTV feeds, FIRs & financial data."),
        ("2. GRAPH MINING", "Extracts call pairs, nocturnal calls (12-6AM) & durations."),
        ("3. CO-LOCATION", "Verifies frequent callers near same CCTV camera zone."),
        ("4. THREAT RATING", "Computes 6-factor composite score (0-100) per suspect."),
        ("5. HASH AUDIT", "Dispatches alert feed & hashes query to SHA-256 chain.")
    ]

    for idx, (stg_title, stg_desc) in enumerate(stages):
        x = Inches(0.8 + idx * 2.38)
        card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(3.3), Inches(2.25), Inches(3.4))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_CYAN if idx == 3 else COLOR_BORDER
        card.line.width = Pt(2) if idx == 3 else Pt(1)

        tf = card.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = stg_title
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_GOLD if idx == 3 else COLOR_CYAN
        p.space_after = Pt(6)

        p_d = tf.add_paragraph()
        p_d.text = stg_desc
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = COLOR_TEXT_MAIN

    # ==========================================
    # SLIDE 4: FEASIBILITY, VIABILITY & PROOF
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_header(slide4, "03/05 [Slide 4]", "Feasibility, Viability & Working Prototype Proof", "Prototype Proof & Feasibility")

    # Left Column: Empirical Test Results
    left_card_s4 = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.2))
    left_card_s4.fill.solid()
    left_card_s4.fill.fore_color.rgb = COLOR_CARD
    left_card_s4.line.color.rgb = COLOR_CYAN
    tf_l4 = left_card_s4.text_frame
    tf_l4.word_wrap = True

    p = tf_l4.paragraphs[0]
    p.text = "EMPIRICAL PROTOTYPE TEST RESULTS"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_GOLD
    p.space_after = Pt(6)

    proto_metrics = [
        ("Total FIR Records Processed:", "121 active police FIR records"),
        ("Call Detail Logs (CDRs):", "182 Call/SMS logs analyzed"),
        ("Suspect Call Interaction Pairs:", "112 suspect pairs mapped in graph"),
        ("CCTV Meetings Confirmed:", "12 physical encounters verified"),
        ("Active Crime Syndicates:", "88 cells detected (Primary: RING-01)"),
        ("Top Threat Identified:", "Md. Ranbir Bhalla (Score: 86.6 / 100)"),
        ("Audit Ledger Verification:", "0 tamper flags (100% SHA-256 Integrity)")
    ]

    for lbl, val in proto_metrics:
        p = tf_l4.add_paragraph()
        p.text = f"• {lbl} "
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_CYAN
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = val
        run.font.bold = False
        run.font.color.rgb = COLOR_TEXT_MAIN

    # Right Column: Deployment Feasibility
    right_card_s4 = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.2))
    right_card_s4.fill.solid()
    right_card_s4.fill.fore_color.rgb = COLOR_CARD
    right_card_s4.line.color.rgb = COLOR_GREEN
    tf_r4 = right_card_s4.text_frame
    tf_r4.word_wrap = True

    p = tf_r4.paragraphs[0]
    p.text = "OPERATIONAL & DEPLOYMENT FEASIBILITY"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_GOLD
    p.space_after = Pt(6)

    feasibility_points = [
        ("Hardware Feasibility:", "Runs on standard laptop/server or ruggedized command van without specialized GPU requirement."),
        ("API Compatibility:", "FastAPI REST API connects seamlessly with state CCTNS databases via REST JSON."),
        ("Container Deployment:", "One-line docker-compose up launch for immediate field command van deployment."),
        ("Sub-Second Latency:", "In-memory Pandas and NetworkX lookup yields < 100ms response times across 10,000+ logs."),
        ("Non-Root Container Security:", "Executes as unprivileged non-root user (UID 1000) with environment key hygiene.")
    ]

    for lbl, val in feasibility_points:
        p = tf_r4.add_paragraph()
        p.text = f"• {lbl} "
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_GREEN
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = val
        run.font.bold = False
        run.font.color.rgb = COLOR_TEXT_MAIN

    # ==========================================
    # SLIDE 5: IMPACT, RISKS & SOCIAL BENEFITS
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_header(slide5, "04/05 [Slide 5]", "Impact, Social Benefits & Potential Risks Mitigation", "Impact & Risk Analysis")

    # 3 Top Banners
    stats = [
        ("80% REDUCTION", "In manual investigation time for mapping crime rings"),
        ("100% AUDITABLE", "Cryptographically verifiable logs for court evidence"),
        ("STATEWIDE SCALE", "Deployable across Maharashtra Police & NCRB")
    ]
    for idx, (num, desc) in enumerate(stats):
        x = Inches(0.8 + idx * 3.95)
        card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.5), Inches(3.7), Inches(1.3))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_CYAN

        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = num
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = COLOR_GOLD
        p.alignment = PP_ALIGN.CENTER

        p_sub = tf.add_paragraph()
        p_sub.text = desc
        p_sub.font.size = Pt(10)
        p_sub.font.color.rgb = COLOR_TEXT_MAIN
        p_sub.alignment = PP_ALIGN.CENTER

    # Left: Social Benefits | Right: Risk Mitigation
    left_impact = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.1), Inches(5.7), Inches(3.6))
    left_impact.fill.solid()
    left_impact.fill.fore_color.rgb = COLOR_CARD
    left_impact.line.color.rgb = COLOR_GREEN
    tf_li = left_impact.text_frame
    tf_li.word_wrap = True

    p = tf_li.paragraphs[0]
    p.text = "LAW ENFORCEMENT & SOCIAL BENEFITS"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_GOLD
    p.space_after = Pt(4)

    bens = [
        "Proactive Policing: Shifts law enforcement from reactive crime investigation to proactive syndicate prevention.",
        "Judicial Integrity: SHA-256 hash chaining eliminates claims of evidence tampering or bias in court.",
        "Resource Optimization: Saves hundreds of officer man-hours spent cross-referencing paper FIRs and call logs."
    ]
    for b in bens:
        p = tf_li.add_paragraph()
        p.text = f"✔ {b}"
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_TEXT_MAIN
        p.space_after = Pt(4)

    right_risk = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(3.1), Inches(5.7), Inches(3.6))
    right_risk.fill.solid()
    right_risk.fill.fore_color.rgb = COLOR_CARD
    right_risk.line.color.rgb = COLOR_RED
    tf_rr = right_risk.text_frame
    tf_rr.word_wrap = True

    p = tf_rr.paragraphs[0]
    p.text = "POTENTIAL RISKS & MITIGATION STRATEGY"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_RED
    p.space_after = Pt(4)

    risks = [
        "Risk: Unauthorized internal data access | Mitigation: Role-Based Access Control & SHA-256 immutable audit logging.",
        "Risk: Data Spoofing or Fake Burner Phones | Mitigation: Multi-modal verification combining CDR with physical CCTV meetings.",
        "Risk: System Downtime in Field Van | Mitigation: Offline Docker container stack with SQLite local fallback."
    ]
    for r in risks:
        p = tf_rr.add_paragraph()
        p.text = f"⚠ {r}"
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_TEXT_MAIN
        p.space_after = Pt(4)

    # ==========================================
    # SLIDE 6: TECH STACK, TEAM & ROADMAP
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6)
    add_header(slide6, "05/05 [Slide 6]", "Technology Stack, Team Capabilities & Future Roadmap", "Tech Stack & Roadmap")

    # Left Box: Tech Stack Quadrants
    left_card_s6 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.2))
    left_card_s6.fill.solid()
    left_card_s6.fill.fore_color.rgb = COLOR_CARD
    left_card_s6.line.color.rgb = COLOR_CYAN
    tf_l6 = left_card_s6.text_frame
    tf_l6.word_wrap = True

    p = tf_l6.paragraphs[0]
    p.text = "TECHNOLOGY STACK SUMMARY"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_GOLD
    p.space_after = Pt(6)

    stack_summary = [
        ("Frontend Dashboard:", "Next.js 14 (App Router), React 18, TypeScript, Tailwind CSS, Leaflet.js, D3.js Force Graphs."),
        ("Backend & Graph Engine:", "Python 3.11, FastAPI REST API, NetworkX Graph Engine, Pandas & NumPy."),
        ("Security & Cryptography:", "SHA-256 Hash Chain Audit, SQLite Ledger, Role-Based Access Control, GET /api/audit/verify."),
        ("DevOps & Containerization:", "Docker, Docker Compose, Non-root User (UID 1000), python-dotenv (.env).")
    ]
    for lbl, val in stack_summary:
        p = tf_l6.add_paragraph()
        p.text = f"• {lbl} "
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_CYAN
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = val
        run.font.bold = False
        run.font.color.rgb = COLOR_TEXT_MAIN

    # Right Top Box: Future Roadmap
    right_top = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.7), Inches(2.7))
    right_top.fill.solid()
    right_top.fill.fore_color.rgb = COLOR_CARD
    right_top.line.color.rgb = COLOR_GOLD
    tf_rt = right_top.text_frame
    tf_rt.word_wrap = True

    p = tf_rt.paragraphs[0]
    p.text = "FUTURE DEVELOPMENT ROADMAP"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_GOLD
    p.space_after = Pt(4)

    roadmap_items = [
        "Phase 1: ANPR (Automatic License Plate Recognition) camera feed integration.",
        "Phase 2: Fine-tuned Local Llama-3 model for automatic unstructured FIR NLP parsing.",
        "Phase 3: Facial Recognition API matching with state criminal mugshot repository.",
        "Phase 4: Encrypted inter-state peer node network for cross-border syndicates."
    ]
    for rm in roadmap_items:
        p = tf_rt.add_paragraph()
        p.text = f"🚀 {rm}"
        p.font.size = Pt(9.5)
        p.font.color.rgb = COLOR_TEXT_MAIN
        p.space_after = Pt(2)

    # Right Bottom Box: Team & Submission Conclusion
    right_bottom = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(4.35), Inches(5.7), Inches(2.35))
    right_bottom.fill.solid()
    right_bottom.fill.fore_color.rgb = COLOR_CARD
    right_bottom.line.color.rgb = COLOR_GREEN
    tf_rb = right_bottom.text_frame
    tf_rb.word_wrap = True

    p = tf_rb.paragraphs[0]
    p.text = "TEAM CAPABILITIES & CONCLUSION"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_GREEN
    p.space_after = Pt(4)

    conc_text = [
        "Team CTRL INNOVATE brings end-to-end expertise in Fullstack Web Development, Network Graph Algorithms, and Security.",
        "Fully operational prototype ready for SIH 2026 evaluators & Brihanmumbai Police Department deployment! 🇮🇳"
    ]
    for ct in conc_text:
        p = tf_rb.add_paragraph()
        p.text = f"🏆 {ct}"
        p.font.size = Pt(9.5)
        p.font.color.rgb = COLOR_TEXT_MAIN
        p.space_after = Pt(4)

    prs.save(output_path)
    print(f"Successfully generated Official SIH 6-Slide Presentation at: {output_path}")

if __name__ == "__main__":
    out_path = os.path.abspath(sys.argv[1]) if len(sys.argv) > 1 else "Brihanmumbai_Police_Tactical_Intelligence_SIH2026.pptx"
    build_sih_6slide_presentation(out_path)
